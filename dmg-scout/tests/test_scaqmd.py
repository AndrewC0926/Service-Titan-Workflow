"""SCAQMD facility grain: XLSX parsing (sheet-picked-by-header, not by
name), full-replace fetch, the normalized-address join onto Ab802Building,
and the normalized-name join used live by ab869_board_rows. See
app/pipeline/scaqmd.py's module docstring for the access investigation
this rests on (FIND and Public Document Search both robots.txt-blocked in
full; this loads South Coast AQMD's own bulk facility-notification XLSX
instead).

The fetch HTTP call is mocked (respx); these tests are about parsing,
full-replace, and the join/ambiguity logic, not live network."""
import io

import httpx
import openpyxl
import pytest
import respx
from sqlmodel import select

from app.http import PoliteClient
from app.models import Ab802Building, HospitalBuilding, ScaqmdFacility, SourceRun, utcnow
from app.pipeline.ab869 import ab869_board_rows
from app.pipeline.scaqmd import (
    AER_XLSX_URL, SOURCE_AER, SOURCE_CARB, _link_ab802, fetch_scaqmd_facilities,
    load_carb_facilities, parse_carb_csv, parse_facility_rows, scaqmd_matches_for_ab869,
)


def fast_client() -> PoliteClient:
    return PoliteClient(interval=0, max_retries=0, respect_robots=False)


HEADER = [
    "Facility ID", "Facility Name ", "Facility  Location Address", " Facility Location City",
    "Facility  Location Zip", "AB 2588", "Criteria Pollutants ≥ 4tpy (100 tpy for CO)",
    "'Core' CTR Facility \n(PTE≥250 tpy, Elevated Toxics, or GHG MRR Facilities)",
    "CTR Phase 3", "Rule 317.1",
]


def _workbook(rows: list[dict], sheet_name="2025 AER", include_notes_sheet=True) -> bytes:
    wb = openpyxl.Workbook()
    if include_notes_sheet:
        notes = wb.active
        notes.title = "Notes"
        notes.append(["Notes:"])
        notes.append([1, "This list is not comprehensive."])
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb.active
        ws.title = sheet_name
    ws.append(HEADER)
    for r in rows:
        ws.append([r.get(h) for h in HEADER])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fac_row(facility_id="41229", name="LUBECO INC", address="123 MAIN ST",
            city="LONG BEACH", zip_code="90802", **flags) -> dict:
    row = {
        "Facility ID": facility_id, "Facility Name ": name,
        "Facility  Location Address": address, " Facility Location City": city,
        "Facility  Location Zip": zip_code,
    }
    for k, v in flags.items():
        row[k] = v
    return row


# ---- parse_facility_rows ---------------------------------------------------

def test_parse_picks_the_data_sheet_by_header_not_by_name():
    """Next year's file has a different sheet name ("2026 AER") -- must
    still parse without a code change."""
    xlsx = _workbook([_fac_row()], sheet_name="2026 AER")
    rows = parse_facility_rows(xlsx)
    assert len(rows) == 1
    assert rows[0]["facility_id"] == "41229"


def test_parse_extracts_all_fields_and_flags():
    xlsx = _workbook([_fac_row(
        facility_id="99999", name="Test Facility", address="500 Elm St", city="Pomona",
        zip_code="91766",
        **{"AB 2588": "X", "Criteria Pollutants ≥ 4tpy (100 tpy for CO)": "X",
           "Rule 317.1": "X"})])
    rows = parse_facility_rows(xlsx)
    r = rows[0]
    assert r["facility_name"] == "Test Facility"
    assert r["address"] == "500 Elm St"
    assert r["city"] == "Pomona"
    assert r["zip_code"] == "91766"
    assert r["ab_2588"] is True
    assert r["meets_ctr_threshold"] is True
    assert r["rule_317_1"] is True
    assert r["core_ctr_facility"] is False  # blank cell -> False, never guessed
    assert r["ctr_phase_3"] is False


def test_parse_drops_rows_with_no_facility_id():
    xlsx = _workbook([_fac_row(facility_id=None), _fac_row(facility_id="1")])
    rows = parse_facility_rows(xlsx)
    assert len(rows) == 1
    assert rows[0]["facility_id"] == "1"


def test_parse_raises_a_clear_error_when_no_matching_sheet_exists():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Not The Right Sheet"
    ws.append(["Something Else Entirely"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="facility id"):
        parse_facility_rows(buf.getvalue())


# ---- fetch_scaqmd_facilities: full-replace, SourceRun, stats ---------------

@respx.mock
def test_fetch_stores_rows_and_a_source_run(db_session, cfg):
    xlsx = _workbook([_fac_row(facility_id="1"), _fac_row(facility_id="2", name="Other Co")])
    respx.get(AER_XLSX_URL).mock(return_value=httpx.Response(200, content=xlsx))

    with fast_client() as client:
        stats = fetch_scaqmd_facilities(db_session, cfg, client)

    assert stats["fetched"] == 2 and stats["stored"] == 2 and stats["error"] is None
    rows = db_session.exec(select(ScaqmdFacility)).all()
    assert {r.facility_id for r in rows} == {"1", "2"}
    assert all(r.in_territory is True for r in rows)

    run = db_session.exec(select(SourceRun)).one()
    assert run.source == "scaqmd_facility" and run.ok is True and run.records_fetched == 2


@respx.mock
def test_fetch_full_replaces_stale_rows_not_seen_in_new_file(db_session, cfg):
    xlsx1 = _workbook([_fac_row(facility_id="1"), _fac_row(facility_id="2")])
    respx.get(AER_XLSX_URL).mock(return_value=httpx.Response(200, content=xlsx1))
    with fast_client() as client:
        fetch_scaqmd_facilities(db_session, cfg, client)
    assert db_session.exec(select(ScaqmdFacility)).all().__len__() == 2

    xlsx2 = _workbook([_fac_row(facility_id="1")])
    respx.get(AER_XLSX_URL).mock(return_value=httpx.Response(200, content=xlsx2))
    with fast_client() as client:
        stats = fetch_scaqmd_facilities(db_session, cfg, client)

    assert stats["stored"] == 1
    remaining = db_session.exec(select(ScaqmdFacility)).all()
    assert {r.facility_id for r in remaining} == {"1"}


@respx.mock
def test_fetch_records_error_on_bad_response_without_raising(db_session, cfg):
    respx.get(AER_XLSX_URL).mock(return_value=httpx.Response(500))
    with fast_client() as client:
        stats = fetch_scaqmd_facilities(db_session, cfg, client)
    assert stats["error"] is not None
    run = db_session.exec(select(SourceRun)).one()
    assert run.ok is False


# ---- _link_ab802: normalized-address match, ambiguity dropped -------------

def _ab802_building(session, property_id, address_1=None, county="Los Angeles",
                    year_ending=2024) -> Ab802Building:
    b = Ab802Building(portfolio_manager_property_id=property_id, address_1=address_1,
                      county_from_geocoding=county, year_ending=year_ending, in_territory=True,
                      source_url="https://example.com")
    session.add(b)
    session.commit()
    return b


def _scaqmd(session, facility_id, address=None, city="Los Angeles") -> ScaqmdFacility:
    f = ScaqmdFacility(facility_id=facility_id, facility_name=f"Facility {facility_id}",
                       address=address, city=city, source_url="https://example.com")
    session.add(f)
    session.commit()
    return f


def test_link_ab802_matches_on_normalized_address(db_session):
    _scaqmd(db_session, "F1", address="123 Main St")
    b = _ab802_building(db_session, "P1", address_1="123 MAIN STREET")  # same street, different suffix spelling

    flagged = _link_ab802(db_session)
    db_session.refresh(b)
    assert flagged == 1
    assert b.air_permit_facility_id == "F1"
    assert b.air_permit_match_method == "normalized_address"


def test_link_ab802_no_match_leaves_field_null(db_session):
    _scaqmd(db_session, "F1", address="123 Main St")
    b = _ab802_building(db_session, "P1", address_1="999 Nowhere Ave")

    _link_ab802(db_session)
    db_session.refresh(b)
    assert b.air_permit_facility_id is None
    assert b.air_permit_match_method is None


def test_link_ab802_ambiguous_scaqmd_side_is_dropped(db_session):
    """Two different SCAQMD facilities sharing one normalized address --
    never guessed at, dropped entirely."""
    _scaqmd(db_session, "F1", address="500 Elm St")
    _scaqmd(db_session, "F2", address="500 Elm St")
    b = _ab802_building(db_session, "P1", address_1="500 Elm St")

    flagged = _link_ab802(db_session)
    db_session.refresh(b)
    assert flagged == 0
    assert b.air_permit_facility_id is None


def test_link_ab802_flags_every_year_of_the_same_building(db_session):
    _scaqmd(db_session, "F1", address="123 Main St")
    b2023 = _ab802_building(db_session, "P1", address_1="123 Main St", year_ending=2023)
    b2024 = _ab802_building(db_session, "P1", address_1="123 Main St", year_ending=2024)

    flagged = _link_ab802(db_session)
    db_session.refresh(b2023)
    db_session.refresh(b2024)
    assert flagged == 2
    assert b2023.air_permit_facility_id == "F1"
    assert b2024.air_permit_facility_id == "F1"


def test_link_ab802_clears_a_stale_match_on_rerun(db_session):
    """A facility dropped from the latest SCAQMD file must un-flag, not
    linger from a prior run."""
    f = _scaqmd(db_session, "F1", address="123 Main St")
    b = _ab802_building(db_session, "P1", address_1="123 Main St")
    _link_ab802(db_session)
    db_session.refresh(b)
    assert b.air_permit_facility_id == "F1"

    db_session.delete(f)
    db_session.commit()
    _link_ab802(db_session)
    db_session.refresh(b)
    assert b.air_permit_facility_id is None


# ---- scaqmd_matches_for_ab869: normalized-name match, same city only ------

def test_ab869_match_on_normalized_name_and_city(db_session):
    f = _scaqmd(db_session, "F1", address="1 Hospital Way", city="Los Angeles")
    f.facility_name = "Test Hospital Inc"
    db_session.add(f)
    db_session.commit()

    matches = scaqmd_matches_for_ab869(db_session, [("PERM1", "Test Hospital, Inc.", "Los Angeles")])
    assert matches == {"PERM1": "F1"}


def test_ab869_match_requires_same_city(db_session):
    f = _scaqmd(db_session, "F1", city="Los Angeles")
    f.facility_name = "Test Hospital"
    db_session.add(f)
    db_session.commit()

    matches = scaqmd_matches_for_ab869(db_session, [("PERM1", "Test Hospital", "Pomona")])
    assert matches == {}


def test_ab869_match_drops_ambiguous_name_city_pair(db_session):
    for fid in ("F1", "F2"):
        f = _scaqmd(db_session, fid, city="Los Angeles")
        f.facility_name = "Community Hospital"
        db_session.add(f)
    db_session.commit()

    matches = scaqmd_matches_for_ab869(db_session, [("PERM1", "Community Hospital", "Los Angeles")])
    assert matches == {}


def test_ab869_match_skips_facility_with_no_city(db_session):
    _scaqmd(db_session, "F1", city=None)
    matches = scaqmd_matches_for_ab869(db_session, [("PERM1", "Some Facility", None)])
    assert matches == {}


# ---- integration: ab869_board_rows carries the flag ------------------------

def test_ab869_board_rows_carries_air_permit_facility_id(db_session, cfg):
    f = _scaqmd(db_session, "F1", city="Los Angeles")
    f.facility_name = "Matched Hospital"
    db_session.add(f)
    db_session.commit()

    hb = HospitalBuilding(perm_id="P1", building_nbr="B1", facility_name="Matched Hospital",
                          city="Los Angeles", county="Los Angeles", state="CA",
                          snapshot_date=utcnow(), source_url="https://example.com")
    db_session.add(hb)
    db_session.commit()

    rows = ab869_board_rows(db_session, cfg)
    row = next(r for r in rows if r["perm_id"] == "P1")
    assert row["air_permit_facility_id"] == "F1"


# ---- CARB CSV: parse_carb_csv, load_carb_facilities -----------------------

CARB_HEADER = ["CO", "AB", "FACID", "DIS", "FNAME", "FSTREET", "FCITY", "FZIP", "FSIC",
              "COID", "DISN", "CHAPIS", "CERR_CODE", "TOGT", "ROGT", "COT", "NOXT", "SOXT",
              "PMT", "PM10T"]


def _carb_csv(rows: list[dict]) -> bytes:
    import csv
    import io as _io

    buf = _io.StringIO()
    w = csv.DictWriter(buf, fieldnames=CARB_HEADER)
    w.writeheader()
    for r in rows:
        w.writerow({h: r.get(h, "") for h in CARB_HEADER})
    return buf.getvalue().encode("utf-8")


def _carb_row(facid="193444", fname="100 N CRESCENT LLC", fstreet="100 N CRESCENT DR",
             fcity="BEVERLY HILLS", fzip="90210") -> dict:
    return {"CO": "19", "AB": "SC", "FACID": facid, "DIS": "SC", "FNAME": fname,
           "FSTREET": fstreet, "FCITY": fcity, "FZIP": fzip}


def test_parse_carb_csv_extracts_fields():
    csv_bytes = _carb_csv([_carb_row()])
    rows = parse_carb_csv(csv_bytes)
    assert len(rows) == 1
    r = rows[0]
    assert r["facility_id"] == "193444"
    assert r["facility_name"] == "100 N CRESCENT LLC"
    assert r["address"] == "100 N CRESCENT DR"
    assert r["city"] == "BEVERLY HILLS"
    assert r["zip_code"] == "90210"


def test_parse_carb_csv_drops_rows_with_no_facid():
    csv_bytes = _carb_csv([_carb_row(facid=""), _carb_row(facid="1")])
    rows = parse_carb_csv(csv_bytes)
    assert len(rows) == 1
    assert rows[0]["facility_id"] == "1"


def test_load_carb_facilities_stores_rows_with_carb_source(db_session, cfg, tmp_path):
    csv_path = tmp_path / "carb.csv"
    csv_path.write_bytes(_carb_csv([_carb_row(facid="1"), _carb_row(facid="2", fname="Other Co")]))

    stats = load_carb_facilities(db_session, csv_path=str(csv_path))
    assert stats["fetched"] == 2 and stats["stored"] == 2 and stats["error"] is None

    rows = db_session.exec(select(ScaqmdFacility).where(ScaqmdFacility.source == SOURCE_CARB)).all()
    assert {r.facility_id for r in rows} == {"1", "2"}
    assert all(r.source == SOURCE_CARB for r in rows)


def test_load_carb_facilities_reports_new_vs_already_present(db_session, cfg, tmp_path):
    # Seed an AER-sourced row under facility_id "1" -- CARB's own row for
    # the same facility_id should count as "already present," not "new".
    db_session.add(ScaqmdFacility(facility_id="1", source=SOURCE_AER, facility_name="AER Facility",
                                  source_url="https://example.com"))
    db_session.commit()

    csv_path = tmp_path / "carb.csv"
    csv_path.write_bytes(_carb_csv([_carb_row(facid="1"), _carb_row(facid="2")]))

    stats = load_carb_facilities(db_session, csv_path=str(csv_path))
    assert stats["already_present"] == 1
    assert stats["new"] == 1


def test_load_carb_facilities_never_deletes_aer_rows(db_session, cfg, tmp_path):
    db_session.add(ScaqmdFacility(facility_id="99", source=SOURCE_AER, facility_name="AER only",
                                  source_url="https://example.com"))
    db_session.commit()

    csv_path = tmp_path / "carb.csv"
    csv_path.write_bytes(_carb_csv([_carb_row(facid="1")]))
    load_carb_facilities(db_session, csv_path=str(csv_path))

    aer_rows = db_session.exec(select(ScaqmdFacility).where(ScaqmdFacility.source == SOURCE_AER)).all()
    assert {r.facility_id for r in aer_rows} == {"99"}


def test_load_carb_facilities_full_replaces_only_carb_rows_on_rerun(db_session, cfg, tmp_path):
    csv_path = tmp_path / "carb.csv"
    csv_path.write_bytes(_carb_csv([_carb_row(facid="1"), _carb_row(facid="2")]))
    load_carb_facilities(db_session, csv_path=str(csv_path))

    csv_path.write_bytes(_carb_csv([_carb_row(facid="1")]))
    stats = load_carb_facilities(db_session, csv_path=str(csv_path))
    assert stats["stored"] == 1

    remaining = db_session.exec(select(ScaqmdFacility).where(ScaqmdFacility.source == SOURCE_CARB)).all()
    assert {r.facility_id for r in remaining} == {"1"}


def test_load_carb_facilities_records_error_on_missing_file(db_session, cfg):
    stats = load_carb_facilities(db_session, csv_path="/nonexistent/path.csv")
    assert stats["error"] is not None
    run = db_session.exec(select(SourceRun).where(SourceRun.source == "scaqmd_facility_carb")).one()
    assert run.ok is False


def test_load_carb_facilities_flags_ab802_the_same_as_aer(db_session, cfg, tmp_path):
    b = _ab802_building(db_session, "P1", address_1="100 N Crescent Dr")
    csv_path = tmp_path / "carb.csv"
    csv_path.write_bytes(_carb_csv([_carb_row(facid="193444", fstreet="100 N Crescent Dr")]))

    stats = load_carb_facilities(db_session, csv_path=str(csv_path))
    db_session.refresh(b)
    assert stats["ab802_flagged"] == 1
    assert b.air_permit_facility_id == "193444"
