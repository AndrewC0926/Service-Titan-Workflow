"""Block 4B Item 6: the contractor handoff list
(app.contractors.buildings_past_service_life_near_contractor and its
xlsx export)."""
from app.contractors import buildings_past_service_life_near_contractor
from app.models import Contractor, RetrofitBuilding


def _contractor(lat=34.00, lon=-118.00):
    return Contractor(license_no="X1", business_name="Test Mechanical", latitude=lat, longitude=lon)


class TestBuildingsPastServiceLifeNearContractor:
    def test_empty_when_contractor_has_no_coordinates(self, db_session):
        contractor = Contractor(license_no="X2", business_name="No Coords Mechanical")
        db_session.add(contractor)
        db_session.commit()
        assert buildings_past_service_life_near_contractor(db_session, contractor, 15) == []

    def test_includes_a_due_building_within_radius(self, db_session):
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(
            apn="pl-1", population="replacement_candidate", address="1 Test Way",
            latitude=34.02, longitude=-118.02, service_life_status="due", year_built=1970,
        ))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15)
        assert len(rows) == 1
        assert rows[0]["address"] == "1 Test Way"
        assert rows[0]["service_life_status"] == "due"
        assert rows[0]["nearest_permit_reference"] is None
        assert rows[0]["equipment_class"] == "unknown"

    def test_excludes_a_building_outside_the_radius(self, db_session):
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(
            apn="pl-2", population="replacement_candidate", address="Far Away",
            latitude=36.00, longitude=-118.00, service_life_status="overdue",
        ))
        db_session.commit()
        assert buildings_past_service_life_near_contractor(db_session, contractor, 15) == []

    def test_excludes_not_due_and_null_service_life_status(self, db_session):
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(
            apn="pl-3", population="replacement_candidate", address="Not Due Yet",
            latitude=34.01, longitude=-118.01, service_life_status="not_due",
        ))
        db_session.add(RetrofitBuilding(
            apn="pl-4", population="replacement_candidate", address="Unknown Status",
            latitude=34.01, longitude=-118.01, service_life_status=None,
        ))
        db_session.commit()
        assert buildings_past_service_life_near_contractor(db_session, contractor, 15) == []

    def test_excludes_recently_active_population(self, db_session):
        """This artifact is specifically "no replacement permit on
        record" -- a recently_active building (has a permit) is a
        different population entirely, never conflated here."""
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(
            apn="pl-5", population="recently_active", address="Has A Permit",
            latitude=34.01, longitude=-118.01, service_life_status="overdue",
            equipment_type="split_dx", latest_install_year=2000,
        ))
        db_session.commit()
        assert buildings_past_service_life_near_contractor(db_session, contractor, 15) == []

    def test_sorted_nearest_first_when_urgency_ties(self, db_session):
        """Distance is only the tie-break, after urgency (years past
        service life, then sqft) -- both null here, so it falls all the
        way through to distance."""
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(apn="pl-6", population="replacement_candidate", address="Far",
                                        latitude=34.10, longitude=-118.10, service_life_status="due"))
        db_session.add(RetrofitBuilding(apn="pl-7", population="replacement_candidate", address="Near",
                                        latitude=34.01, longitude=-118.01, service_life_status="due"))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15)
        assert [r["address"] for r in rows] == ["Near", "Far"]

    def test_sorted_by_years_past_service_life_first(self, db_session):
        """Block 4C: urgency (years past service life desc) outranks
        distance -- a farther, more-overdue building comes first."""
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(apn="pl-9", population="replacement_candidate", address="Near, mild",
                                        latitude=34.01, longitude=-118.01, service_life_status="due",
                                        service_life_years_past=2.0))
        db_session.add(RetrofitBuilding(apn="pl-10", population="replacement_candidate", address="Far, severe",
                                        latitude=34.10, longitude=-118.10, service_life_status="overdue",
                                        service_life_years_past=40.0))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15)
        assert [r["address"] for r in rows] == ["Far, severe", "Near, mild"]

    def test_sqft_breaks_a_years_past_tie(self, db_session):
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(apn="pl-11", population="replacement_candidate", address="Small",
                                        latitude=34.01, longitude=-118.01, service_life_status="due",
                                        service_life_years_past=10.0, sqft=5000))
        db_session.add(RetrofitBuilding(apn="pl-12", population="replacement_candidate", address="Big",
                                        latitude=34.02, longitude=-118.02, service_life_status="due",
                                        service_life_years_past=10.0, sqft=50000))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15)
        assert [r["address"] for r in rows] == ["Big", "Small"]

    def test_defaults_to_top_40_by_urgency(self, db_session):
        contractor = _contractor()
        db_session.add(contractor)
        for i in range(50):
            db_session.add(RetrofitBuilding(
                apn=f"lim-{i}", population="replacement_candidate", address=f"Building {i}",
                latitude=34.01, longitude=-118.01, service_life_status="overdue",
                service_life_years_past=float(i),
            ))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15)
        assert len(rows) == 40
        # highest years-past-service-life first -- i=49 down to i=10
        assert rows[0]["address"] == "Building 49"
        assert rows[-1]["address"] == "Building 10"

    def test_limit_none_returns_every_matching_row(self, db_session):
        contractor = _contractor()
        db_session.add(contractor)
        for i in range(50):
            db_session.add(RetrofitBuilding(
                apn=f"nolim-{i}", population="replacement_candidate", address=f"Building {i}",
                latitude=34.01, longitude=-118.01, service_life_status="overdue",
                service_life_years_past=float(i),
            ))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15, limit=None)
        assert len(rows) == 50

    def test_limit_100_returns_up_to_100(self, db_session):
        contractor = _contractor()
        db_session.add(contractor)
        for i in range(50):
            db_session.add(RetrofitBuilding(
                apn=f"lim100-{i}", population="replacement_candidate", address=f"Building {i}",
                latitude=34.01, longitude=-118.01, service_life_status="overdue",
                service_life_years_past=float(i),
            ))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15, limit=100)
        assert len(rows) == 50  # fewer than the limit exist, all returned

    def test_uses_real_equipment_class_when_somehow_present(self, db_session):
        """Not the expected shape for this population (equipment_type is
        always null for replacement_candidate rows by construction), but
        the function must still classify it correctly rather than assume
        null."""
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(
            apn="pl-8", population="replacement_candidate", address="Edge Case",
            latitude=34.01, longitude=-118.01, service_life_status="due", equipment_type="boiler",
        ))
        db_session.commit()
        rows = buildings_past_service_life_near_contractor(db_session, contractor, 15)
        assert rows[0]["equipment_class"] == "boiler"


class TestContractorBuildingsXlsxRoute:
    def _client(self, db_session, monkeypatch):
        import base64

        from fastapi.testclient import TestClient

        from app.db import get_session
        from app.web.main import app

        monkeypatch.setenv("DASHBOARD_PASSWORD", "testpw")
        app.dependency_overrides[get_session] = lambda: db_session
        client = TestClient(app)
        auth = {"Authorization": "Basic " + base64.b64encode(b"andrew:testpw").decode()}
        return client, auth, app

    def test_downloads_a_real_xlsx_with_expected_rows(self, db_session, monkeypatch):
        contractor = _contractor()
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(
            apn="route-1", population="replacement_candidate", address="1 Route St",
            latitude=34.02, longitude=-118.02, service_life_status="overdue",
        ))
        db_session.commit()

        client, auth, app = self._client(db_session, monkeypatch)
        resp = client.get(f"/contractor/{contractor.id}/buildings.xlsx", headers=auth)
        app.dependency_overrides.clear()
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        from io import BytesIO

        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(resp.content))
        ws = wb["Buildings"]
        assert ws.cell(row=3, column=1).value == "1 Route St"

    def test_filename_carries_contractor_name_and_date(self, db_session, monkeypatch):
        from datetime import date

        contractor = _contractor()
        contractor.business_name = "Test Mechanical, Inc."
        db_session.add(contractor)
        db_session.add(RetrofitBuilding(
            apn="fn-1", population="replacement_candidate", address="1 Route St",
            latitude=34.02, longitude=-118.02, service_life_status="overdue",
        ))
        db_session.commit()

        client, auth, app = self._client(db_session, monkeypatch)
        resp = client.get(f"/contractor/{contractor.id}/buildings.xlsx", headers=auth)
        app.dependency_overrides.clear()
        assert resp.status_code == 200
        disposition = resp.headers["content-disposition"]
        assert "Test-Mechanical-Inc" in disposition
        assert date.today().isoformat() in disposition

    def test_limit_query_param_is_honored(self, db_session, monkeypatch):
        from io import BytesIO

        from openpyxl import load_workbook

        contractor = _contractor()
        db_session.add(contractor)
        for i in range(50):
            db_session.add(RetrofitBuilding(
                apn=f"rt-{i}", population="replacement_candidate", address=f"Building {i}",
                latitude=34.01, longitude=-118.01, service_life_status="overdue",
                service_life_years_past=float(i),
            ))
        db_session.commit()

        client, auth, app = self._client(db_session, monkeypatch)
        default_resp = client.get(f"/contractor/{contractor.id}/buildings.xlsx", headers=auth)
        limited_resp = client.get(f"/contractor/{contractor.id}/buildings.xlsx?limit=100", headers=auth)
        app.dependency_overrides.clear()

        default_wb = load_workbook(BytesIO(default_resp.content))
        limited_wb = load_workbook(BytesIO(limited_resp.content))
        # header row (1) + column-title row (2) + N data rows
        assert default_wb["Buildings"].max_row == 2 + 40
        assert limited_wb["Buildings"].max_row == 2 + 50

    def test_404_for_an_unknown_contractor(self, db_session, monkeypatch):
        client, auth, app = self._client(db_session, monkeypatch)
        resp = client.get("/contractor/999999/buildings.xlsx", headers=auth)
        app.dependency_overrides.clear()
        assert resp.status_code == 404

    def test_link_appears_on_the_accounts_contractors_tab(self, db_session, monkeypatch):
        contractor = _contractor()
        contractor.classifications = "C20"  # the default "mechanical" filter requires this
        db_session.add(contractor)
        db_session.commit()
        client, auth, app = self._client(db_session, monkeypatch)
        resp = client.get("/accounts/contractors", headers=auth)
        app.dependency_overrides.clear()
        assert resp.status_code == 200
        assert f"/contractor/{contractor.id}/buildings.xlsx" in resp.text

    def test_link_appears_on_a_contractor_anchored_pipeline_row(self, db_session, monkeypatch):
        from app.models import Contact, Opportunity, PenState, ReasonBlock, ReasonStrength, Signal, SignalType, WhyKind

        contractor = _contractor()
        db_session.add(contractor)
        db_session.flush()
        contact = Contact(name="Someone", phone="555-1111", reach_status="confirmed",
                          contractor_id=contractor.id)
        signal = Signal(signal_type=SignalType.ceqa_nop)
        db_session.add(contact)
        db_session.add(signal)
        db_session.flush()
        opp = Opportunity(account_id=1, contact_id=contact.id, signal_id=signal.id,
                          pen_state=PenState.not_moved, owner_user="andrew")
        db_session.add(opp)
        db_session.flush()
        for kind in WhyKind:
            db_session.add(ReasonBlock(opportunity_id=opp.id, why_kind=kind,
                                       strength=ReasonStrength.Strong, evidence="x"))
        db_session.commit()

        client, auth, app = self._client(db_session, monkeypatch)
        resp = client.get("/pipeline", headers=auth)
        app.dependency_overrides.clear()
        assert resp.status_code == 200
        assert f"/contractor/{contractor.id}/buildings.xlsx" in resp.text

    def test_no_link_for_a_non_contractor_anchored_pipeline_row(self, db_session, monkeypatch):
        from app.models import Contact, Opportunity, PenState, ReasonBlock, ReasonStrength, Signal, SignalType, WhyKind

        contact = Contact(name="Project Contact", phone="555-2222", reach_status="confirmed")
        signal = Signal(signal_type=SignalType.ceqa_nop)
        db_session.add(contact)
        db_session.add(signal)
        db_session.flush()
        opp = Opportunity(account_id=1, contact_id=contact.id, signal_id=signal.id,
                          pen_state=PenState.not_moved, owner_user="andrew")
        db_session.add(opp)
        db_session.flush()
        for kind in WhyKind:
            db_session.add(ReasonBlock(opportunity_id=opp.id, why_kind=kind,
                                       strength=ReasonStrength.Strong, evidence="x"))
        db_session.commit()

        client, auth, app = self._client(db_session, monkeypatch)
        resp = client.get("/pipeline", headers=auth)
        app.dependency_overrides.clear()
        assert resp.status_code == 200
        assert "buildings.xlsx" not in resp.text
