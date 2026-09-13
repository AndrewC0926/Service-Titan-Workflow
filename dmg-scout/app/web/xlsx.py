"""Block 4B Item 3 (Master Plan v3.6 section 40): "xlsx export of the
underlying rows" for the Weekly Sales Intelligence Brief. One sheet per
section, plain rows -- a rep or Larry pastes this into whatever they
already use, not a second dashboard.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook


def weekly_brief_xlsx(brief: dict) -> bytes:
    wb = Workbook()

    funnel_ws = wb.active
    funnel_ws.title = "Funnel"
    funnel_ws.append(["Stage", "Value", "vs last week"])
    for f in brief["funnel"]:
        funnel_ws.append([f["label"], f["value"], f["wow"]["label"] if f["wow"] else "not enough history"])

    made_ws = wb.create_sheet("Moves made")
    made_ws.append(["When", "Who", "What"])
    for m in brief["moves_made"]:
        made_ws.append([m["at"], m.get("user") or m.get("author"), m["text"]])

    next_ws = wb.create_sheet("Moves next")
    next_ws.append(["Opportunity", "Account or building", "Weakest why"])
    for r in brief["moves_next"]:
        next_ws.append([r["opportunity_id"], r["account_or_building"], r["weakest_why"] or "unknown, not guessed"])

    deadline_ws = wb.create_sheet("New deadline exposure")
    deadline_ws.append(["Regulation", "Account or building", "Date"])
    for row in brief["new_deadline_exposure"]["rows"]:
        deadline_ws.append([row["regulation"], row["account_or_building"], row["date"]])

    lead_ws = wb.create_sheet("One lead")
    lead_ws.append(["Why", "Strength", "Evidence"])
    if brief["one_lead"]:
        for rb in brief["one_lead"]["reason_blocks"]:
            lead_ws.append([rb["why_kind"], rb["strength"], rb["evidence"]])
        if brief["one_lead_do"]:
            lead_ws.append(["do", "", brief["one_lead_do"]])

    wrong_ws = wb.create_sheet("What Scout got wrong")
    wrong_ws.append(["Opportunity", "User", "When"])
    for w in brief["got_wrong"]:
        wrong_ws.append([w["opportunity_id"], w["user"], w["created_at"]])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
