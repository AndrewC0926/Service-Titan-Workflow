"""AB 802 statewide benchmarking import -- energy.ca.gov's annual "Download
submitted {year} benchmarking information" file. See app/models.py's
Ab802Building docstring for the full design (keying, full-replace-per-year,
no owner column, the loose assessor/EBEWE join).

COMPLIANCE, checked before this shipped (Phase A research, 2026-09-06):
energy.ca.gov/robots.txt disallows /admin/, /comment/reply/, /search/,
/user/{register,password,login,logout}, /media/oembed -- NOT /media/{id}
landing pages and NOT /sites/default/files/, where the actual xlsx lives.
Site-wide terms are the same ca.gov Conditions of Use CAEATFA's own module
docstring already quotes (public domain, "may be distributed or copied as
permitted by law", no automated-access restriction) -- the same statewide
CDT policy, not a separate finding per department. See app/assumptions.py's
"AB 802 statewide benchmarking" group for the recorded classification.

FETCH CADENCE: annual, not daily. The file is published once a year with no
fixed date (the 2024 file's Last-Modified is 2025-10-08; both the 2023 and
2024 files were touched together that day -- an ADA-remediation republish,
not a new data pull, since the 2023 file's own header row says "downloaded
July 16, 2025" despite an October Last-Modified). NOT called from `scout
pipeline` (app.cli.pipeline) -- a human runs `scout fetch-ab802-benchmarks
<year>` by hand, same "off the daily/weekly schedule" treatment
app/pipeline/hcai.py's manual imports get, except this one really is an
automated fetcher, just not on an automatic cadence.

COLUMN LAYOUT: positions, not header-text matching -- same discipline
app/sources/caeatfa.py's COL_* constants already use, and for the same
reason: matching on exact punctuation (this file uses "²" and an en dash in
several headers) is more fragile than a fixed position, verified against a
live download and re-checked here at parse time rather than trusted
silently forever.

JOIN METHOD: lat/long first (RetrofitBuilding.latitude/longitude, itself
rejoined from RetrofitGeocode -- see that model's docstring), within
MATCH_RADIUS_METERS; falls back to app.pipeline.retrofit:normalize_address
text matching. Unlike EBEWE's own join (app/pipeline/ebewe.py's module
docstring: independent Census geocoding on BOTH sides came out WORSE than a
plain text join for this exact LA-parcel-density problem), neither side
here needs a NEW geocoding run: AB 802's own Portfolio Manager geocoding is
already in the file, and RetrofitGeocode's Census geocoding is already
built and stored on RetrofitBuilding -- so there is no fresh jitter being
introduced, and lat/long is tried first. Same abstain-don't-guess ambiguity
discipline as the EBEWE join throughout: a match claimed by more than one
row on either side is dropped, never guessed at. See app/assumptions.py for
the measured match rate and method breakdown.
"""
from __future__ import annotations

import io
import logging
import re
import statistics

import openpyxl
from sqlmodel import Session, delete, select

from app.config import Config
from app.contractors import haversine_miles
from app.http import PoliteClient
from app.models import Ab802Building, EbeweBenchmark, RetrofitBuilding, SourceRun, utcnow
from app.pipeline.retrofit import normalize_address
from app.pipeline.size_score import in_territory as _territory_check

log = logging.getLogger(__name__)

SOURCE = "ab802_benchmarking"
FILE_URL_TEMPLATE = "https://www.energy.ca.gov/sites/default/files/2025-10/{year}_Download_ADA.xlsx"
PORTAL_URL = "https://www.energy.ca.gov/programs-and-topics/programs/building-energy-benchmarking-program"
HEADER_ROW = 3       # 1-indexed -- row 1 is a title line, row 2 is blank; data starts row 4
EXPECTED_MIN_COLUMNS = 34   # verified live against the 2023 and 2024 files, 2026-09-06

# 0-indexed column positions -- see module docstring on why position, not text.
COL_PROPERTY_ID = 0
COL_STANDARD_ID = 1
COL_PROPERTY_NAME = 2
COL_ADDRESS_1 = 3
COL_CITY = 4
COL_STATE = 5
COL_POSTAL_CODE = 6
COL_GFA = 7
COL_PROPERTY_TYPE = 8
COL_ALL_USE_TYPES = 9
COL_SITE_EUI = 10
COL_NATURAL_GAS = 11
COL_ELECTRICITY_GRID = 12
COL_ELECTRICITY_ONSITE = 13
COL_FUEL_OIL_2 = 14
COL_DISTRICT_STEAM = 15
COL_DIESEL = 16
COL_PROPANE = 17
COL_DISTRICT_HOT_WATER = 18
COL_DISTRICT_CHILLED_WATER = 19
COL_YEAR_BUILT = 20
COL_W_ENERGY = 21
COL_ENERGY_STAR_SCORE = 22
COL_ENERGY_STAR_CERTIFIED = 23
COL_ENERGY_STAR_CERT_YEARS = 24
COL_LATITUDE = 25
COL_LONGITUDE = 26
COL_COUNTY = 27
COL_REPORT_GEN_DATE = 28
COL_USED_ESTIMATED = 29
COL_ALERT_PARTIAL_YEAR = 30
COL_GHG_TOTAL = 31
COL_GHG_INTENSITY = 32
COL_YEAR_ENDING = 33

MATCH_RADIUS_METERS = 30.0
_METERS_PER_MILE = 1609.344

_NOT_AVAILABLE = {"", "not available", "na", "n/a"}


def _num(val) -> float | None:
    if val is None:
        return None
    s = str(val).strip()
    if s.lower() in _NOT_AVAILABLE:
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def _int(val) -> int | None:
    n = _num(val)
    return int(n) if n is not None else None


def _str(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return s if s and s.lower() not in _NOT_AVAILABLE else None


def _cell(row: tuple, idx: int):
    return row[idx] if idx < len(row) else None


def parse_rows(raw_bytes: bytes, year: int) -> list[dict]:
    """One dict per data row, field names matching Ab802Building 1:1.
    `year` is used only as the Year Ending fallback for a row that somehow
    lacks its own (never observed live, but the file has no other source
    of "which annual file did this come from")."""
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
    header = next(rows_iter)
    if len(header) < EXPECTED_MIN_COLUMNS:
        log.warning("AB 802 %s file header has %d columns, expected >= %d -- "
                    "layout may have changed, verify COL_* constants against a live download",
                    year, len(header), EXPECTED_MIN_COLUMNS)

    out = []
    for row in rows_iter:
        if all(v is None for v in row):
            continue
        property_id = _str(_cell(row, COL_PROPERTY_ID))
        if not property_id:
            continue  # unkeyable row -- can't upsert without the property id half of the unique key
        out.append({
            "portfolio_manager_property_id": property_id,
            "standard_id": _str(_cell(row, COL_STANDARD_ID)),
            "property_name": _str(_cell(row, COL_PROPERTY_NAME)),
            "address_1": _str(_cell(row, COL_ADDRESS_1)),
            "city": _str(_cell(row, COL_CITY)),
            "state_province": _str(_cell(row, COL_STATE)),
            "postal_code": _str(_cell(row, COL_POSTAL_CODE)),
            "property_gfa_sqft": _num(_cell(row, COL_GFA)),
            "primary_property_type": _str(_cell(row, COL_PROPERTY_TYPE)),
            "all_property_use_types": _str(_cell(row, COL_ALL_USE_TYPES)),
            "weather_normalized_site_eui": _num(_cell(row, COL_SITE_EUI)),
            "natural_gas_use_kbtu": _num(_cell(row, COL_NATURAL_GAS)),
            "electricity_grid_purchase_kbtu": _num(_cell(row, COL_ELECTRICITY_GRID)),
            "electricity_onsite_renewable_kbtu": _num(_cell(row, COL_ELECTRICITY_ONSITE)),
            "fuel_oil_2_use_kbtu": _num(_cell(row, COL_FUEL_OIL_2)),
            "district_steam_use_kbtu": _num(_cell(row, COL_DISTRICT_STEAM)),
            "diesel_use_kbtu": _num(_cell(row, COL_DIESEL)),
            "propane_use_kbtu": _num(_cell(row, COL_PROPANE)),
            "district_hot_water_use_kbtu": _num(_cell(row, COL_DISTRICT_HOT_WATER)),
            "district_chilled_water_use_kbtu": _num(_cell(row, COL_DISTRICT_CHILLED_WATER)),
            "year_built": _int(_cell(row, COL_YEAR_BUILT)),
            "w_energy": _str(_cell(row, COL_W_ENERGY)),
            "energy_star_score": _int(_cell(row, COL_ENERGY_STAR_SCORE)),
            "energy_star_certified": _str(_cell(row, COL_ENERGY_STAR_CERTIFIED)),
            "energy_star_cert_years": _str(_cell(row, COL_ENERGY_STAR_CERT_YEARS)),
            "latitude": _num(_cell(row, COL_LATITUDE)),
            "longitude": _num(_cell(row, COL_LONGITUDE)),
            "county_from_geocoding": _str(_cell(row, COL_COUNTY)),
            "report_generation_date": _str(_cell(row, COL_REPORT_GEN_DATE)),
            "used_estimated_energy_values": _str(_cell(row, COL_USED_ESTIMATED)),
            "alert_partial_year_data": _str(_cell(row, COL_ALERT_PARTIAL_YEAR)),
            "total_ghg_emissions_metric_tons": _num(_cell(row, COL_GHG_TOTAL)),
            "ghg_emissions_intensity": _num(_cell(row, COL_GHG_INTENSITY)),
            "year_ending": _int(_cell(row, COL_YEAR_ENDING)) or year,
        })
    wb.close()
    return out


def fetch_ab802_benchmarks(session: Session, cfg: Config, client: PoliteClient, year: int) -> dict:
    """Downloads one year's file, full-replaces every Ab802Building row for
    that year_ending, computes in_territory, then joins the fresh rows
    against RetrofitBuilding and EbeweBenchmark. Safe to re-run: other
    years already on file are untouched (only `year_ending == year` rows
    are deleted first).

    DUPLICATE PROPERTY IDS, confirmed live against both the 2023 and 2024
    files (2026-09-07): the file itself is not unique per Portfolio Manager
    Property ID within a year -- 50 of 24,758 rows in 2023, 39 of 25,591 in
    2024, almost always the same property resubmitted with a corrected
    address (e.g. "lombard street" vs "1 lombard street" for the same
    building), occasionally two genuinely different addresses under one
    id. The LAST row for a given id, in the file's own row order, is kept
    -- CEC's own file lists a correction after the original it corrects,
    never before, so this is the file's own stated sequence, not a guess
    at which row is "right." The dropped count is returned in stats
    (`duplicate_property_ids_dropped`) rather than silently discarded."""
    url = FILE_URL_TEMPLATE.format(year=year)
    run = SourceRun(source=SOURCE)
    session.add(run)
    session.commit()

    fetched, stored, duplicates_dropped, error = 0, 0, 0, None
    try:
        raw = client.get_bytes(url)
        rows = parse_rows(raw, year)
        fetched = len(rows)

        deduped: dict[str, dict] = {}
        for parsed in rows:
            deduped[parsed["portfolio_manager_property_id"]] = parsed  # last one in file order wins
        duplicates_dropped = fetched - len(deduped)
        if duplicates_dropped:
            log.warning("AB 802 %s: %d rows shared a Portfolio Manager Property ID with a later row "
                       "-- keeping the last one per id, dropping %d", year, duplicates_dropped, duplicates_dropped)

        session.exec(delete(Ab802Building).where(Ab802Building.year_ending == year))

        territory_cfg = cfg.get("territory") or {}
        new_rows: list[Ab802Building] = []
        for parsed in deduped.values():
            county = parsed["county_from_geocoding"]
            state = parsed["state_province"]
            territory = bool(territory_cfg) and _territory_check(cfg, state, county)
            rec = Ab802Building(source_url=url, in_territory=territory, **parsed)
            session.add(rec)
            new_rows.append(rec)
        session.flush()
        stored = len(new_rows)

        la_rows = [r for r in new_rows if r.in_territory and _is_la_county(r.county_from_geocoding)]
        match_to_retrofit(session, la_rows)
        match_to_ebewe(session, la_rows)

        session.commit()
    except Exception as exc:  # noqa: BLE001 -- recorded on the SourceRun, not raised past this stage
        session.rollback()
        error = f"{type(exc).__name__}: {exc}"
        log.error("AB 802 %s fetch failed: %s", year, error)

    run.finished_at = utcnow()
    run.records_fetched = fetched
    run.records_new = stored
    run.ok = error is None
    run.error = error
    session.add(run)
    session.commit()
    return {"year": year, "fetched": fetched, "stored": stored,
           "duplicate_property_ids_dropped": duplicates_dropped, "error": error}


def _is_la_county(county: str | None) -> bool:
    from app.normalize import normalize_county
    return normalize_county(county) == "Los Angeles"


_GRID_DEGREES = 0.001  # ~111m of latitude, ~90m of longitude at LA's latitude --
                       # comfortably larger than MATCH_RADIUS_METERS, so a point's own
                       # cell plus its 8 neighbors always covers the full 30m radius.


def _grid_cell(lat: float, lon: float) -> tuple[int, int]:
    return (round(lat / _GRID_DEGREES), round(lon / _GRID_DEGREES))


def _grid_neighbors(lat: float, lon: float) -> list[tuple[int, int]]:
    gy, gx = _grid_cell(lat, lon)
    return [(gy + dy, gx + dx) for dy in (-1, 0, 1) for dx in (-1, 0, 1)]


def match_to_retrofit(session: Session, rows: list[Ab802Building]) -> int:
    """Lat/long first (within MATCH_RADIUS_METERS), falling back to
    normalized-address text matching -- see module docstring. Ambiguity is
    excluded in both directions, same discipline as
    app.pipeline.ebewe:ebewe_matches_by_normalized_address: a retrofit
    building claimed by more than one AB 802 row (or vice versa) is
    dropped, not guessed at. Writes assessor_match_method/
    assessor_match_distance_m/retrofit_apn directly onto each row in
    `rows`. Returns the number matched."""
    if not rows:
        return 0

    geocoded_rb = session.exec(
        select(RetrofitBuilding.apn, RetrofitBuilding.latitude, RetrofitBuilding.longitude,
              RetrofitBuilding.address)
        .where(RetrofitBuilding.latitude.is_not(None), RetrofitBuilding.longitude.is_not(None))
    ).all()
    # Grid-indexed, same shape as the geocoder comparison app/pipeline/ebewe.py's
    # docstring describes -- avoids an O(len(rows) * len(geocoded_rb)) scan
    # against every LA-County retrofit building for every AB 802 row.
    grid: dict[tuple[int, int], list[tuple[str, float, float]]] = {}
    for apn, rb_lat, rb_lon, _addr in geocoded_rb:
        grid.setdefault(_grid_cell(rb_lat, rb_lon), []).append((apn, rb_lat, rb_lon))

    matched = 0
    # latlong pass -- collect every (row, apn, distance) candidate within
    # radius first, so both-direction ambiguity can be checked before any
    # row is actually assigned a match.
    candidates: list[tuple[Ab802Building, str, float]] = []
    for r in rows:
        if r.latitude is None or r.longitude is None:
            continue
        seen_apns: set[str] = set()
        for cell in _grid_neighbors(r.latitude, r.longitude):
            for apn, rb_lat, rb_lon in grid.get(cell, []):
                if apn in seen_apns:
                    continue
                seen_apns.add(apn)
                d_meters = haversine_miles(r.latitude, r.longitude, rb_lat, rb_lon) * _METERS_PER_MILE
                if d_meters <= MATCH_RADIUS_METERS:
                    candidates.append((r, apn, d_meters))

    claims_by_row: dict[int, list[tuple[str, float]]] = {}
    claims_by_apn: dict[str, list[int]] = {}
    for r, apn, dist in candidates:
        claims_by_row.setdefault(id(r), []).append((apn, dist))
        claims_by_apn.setdefault(apn, []).append(id(r))

    latlong_matched_row_ids: set[int] = set()
    for r, apn, dist in candidates:
        if len(claims_by_row[id(r)]) > 1:
            continue  # this row is within radius of more than one apn -- ambiguous, abstain
        if len(claims_by_apn[apn]) > 1:
            continue  # this apn is within radius of more than one row -- ambiguous, abstain
        r.assessor_match_method = "latlong"
        r.assessor_match_distance_m = round(dist, 1)
        r.retrofit_apn = apn
        latlong_matched_row_ids.add(id(r))
        matched += 1

    # normalized_address fallback, for rows the latlong pass didn't resolve
    remaining = [r for r in rows if id(r) not in latlong_matched_row_ids]
    if remaining:
        rb_by_addr: dict[str, list[str]] = {}
        for apn, _lat, _lon, address in geocoded_rb:
            norm = normalize_address(address)
            if norm:
                rb_by_addr.setdefault(norm, []).append(apn)
        # also cover retrofit rows with no coordinates at all -- an address
        # match doesn't need one
        no_coord_rb = session.exec(
            select(RetrofitBuilding.apn, RetrofitBuilding.address)
            .where(RetrofitBuilding.latitude.is_(None))
        ).all()
        for apn, address in no_coord_rb:
            norm = normalize_address(address)
            if norm:
                rb_by_addr.setdefault(norm, []).append(apn)
        unambiguous_rb = {addr: apns[0] for addr, apns in rb_by_addr.items() if len(apns) == 1}

        addr_by_row: dict[str, list[Ab802Building]] = {}
        for r in remaining:
            norm = normalize_address(r.address_1)
            if norm:
                addr_by_row.setdefault(norm, []).append(r)

        for addr, ab_rows in addr_by_row.items():
            if len(ab_rows) > 1:
                continue  # more than one AB 802 row shares this address -- ambiguous, abstain
            apn = unambiguous_rb.get(addr)
            if apn is None:
                continue
            r = ab_rows[0]
            r.assessor_match_method = "normalized_address"
            r.assessor_match_distance_m = None
            r.retrofit_apn = apn
            matched += 1

    return matched


def match_to_ebewe(session: Session, rows: list[Ab802Building]) -> int:
    """Normalized-address match against EbeweBenchmark (no lat/long side on
    that table at all) -- most-recent program_year per building_id, same
    both-direction ambiguity exclusion as match_to_retrofit. Writes
    ebewe_building_id/benchmarking_filer -- see Ab802Building's docstring
    for why that field is never called "owner"."""
    if not rows:
        return 0

    latest: dict[str, EbeweBenchmark] = {}
    for rec in session.exec(select(EbeweBenchmark)).all():
        cur = latest.get(rec.building_id)
        if cur is None or rec.program_year > cur.program_year:
            latest[rec.building_id] = rec

    ebewe_by_addr: dict[str, list[EbeweBenchmark]] = {}
    for rec in latest.values():
        norm = normalize_address(rec.building_address)
        if norm:
            ebewe_by_addr.setdefault(norm, []).append(rec)
    unambiguous_ebewe = {addr: recs[0] for addr, recs in ebewe_by_addr.items() if len(recs) == 1}

    addr_by_row: dict[str, list[Ab802Building]] = {}
    for r in rows:
        norm = normalize_address(r.address_1)
        if norm:
            addr_by_row.setdefault(norm, []).append(r)

    matched = 0
    for addr, ab_rows in addr_by_row.items():
        if len(ab_rows) > 1:
            continue
        rec = unambiguous_ebewe.get(addr)
        if rec is None or not rec.organization:
            continue
        r = ab_rows[0]
        r.ebewe_building_id = rec.building_id
        r.benchmarking_filer = rec.organization
        matched += 1

    return matched


# ---- board: latest-year snapshot, ranked, filtered ------------------------


# Age-credit trapezoid -- see age_credit()'s own docstring for the reasoning.
AGE_CREDIT_ZERO_BEFORE = 1975   # pre-1975: no credit at all, not full credit
AGE_CREDIT_FULL_START = 1985    # ramps 1975 -> 1985 from 0 to 1.0
AGE_CREDIT_FULL_END = 2012      # full credit (1.0) for 1985-2012 inclusive
AGE_CREDIT_ZERO_AFTER = 2018    # ramps 2012 -> 2018 from 1.0 back to 0


def age_credit(year_built: int | None) -> float | None:
    """0..1: full credit (1.0) for Year Built 1985-2012 inclusive, tapering
    LINEARLY to 0 by 1975 on the early side and by 2018 on the late side.
    None (year_built itself unstated) returns None, never 0 -- 0 would look
    identical to "confirmed too new or too old for credit," a different
    fact than "not stated at all."

    Pre-1975 is 0, not full credit, even though it's the oldest building
    stock on the board: by now, a pre-1975 building's ORIGINAL equipment
    has almost certainly already been replaced at least once (typical
    commercial HVAC service life is 15-25 years), so Year Built alone
    tells you nothing about the age of whatever equipment is actually in
    the building today. That is genuinely unknown, not "assumed old" --
    same abstain-don't-guess discipline as every other null-shaped fact in
    this codebase, just expressed as a credit of 0 instead of a null field,
    since the age term must always produce a number to sum with the EUI
    term below."""
    if year_built is None:
        return None
    if year_built < AGE_CREDIT_ZERO_BEFORE or year_built >= AGE_CREDIT_ZERO_AFTER:
        return 0.0
    if AGE_CREDIT_FULL_START <= year_built <= AGE_CREDIT_FULL_END:
        return 1.0
    if year_built < AGE_CREDIT_FULL_START:
        return (year_built - AGE_CREDIT_ZERO_BEFORE) / (AGE_CREDIT_FULL_START - AGE_CREDIT_ZERO_BEFORE)
    return (AGE_CREDIT_ZERO_AFTER - year_built) / (AGE_CREDIT_ZERO_AFTER - AGE_CREDIT_FULL_END)


def latest_in_territory_rows(session: Session) -> list[Ab802Building]:
    """One row per building -- the most recent year_ending on file for each
    portfolio_manager_property_id, in-territory only. A board snapshot, not
    the year-over-year record: query Ab802Building directly, filtered by
    portfolio_manager_property_id, for that."""
    all_rows = session.exec(
        select(Ab802Building).where(Ab802Building.in_territory == True)  # noqa: E712
        .order_by(Ab802Building.year_ending.desc())
    ).all()
    latest: dict[str, Ab802Building] = {}
    for r in all_rows:
        latest.setdefault(r.portfolio_manager_property_id, r)
    return list(latest.values())


# Default property-type filter -- see rank_in_territory's `restrict_to_relevant_types`.
# Verified 2026-09-07 against the exact primary_property_type strings live in
# production ab802_buildings (2024, in-territory) -- an ALLOWLIST, not an
# exclude list: golf/country club rows carry "Other - Recreation" or "Other -
# Restaurant/Bar", casino rows carry "Casino", worship rows "Worship Facility",
# parking rows "Parking" -- none of those, and no other "Other"/"Other - X"
# catch-all (Portfolio Manager's own miscellaneous bucket, never a specific
# type DMG would target with confidence), are in this set, so they're excluded
# by construction rather than by naming every excluded type. Self-Storage
# Facility, Senior Living Community, and Residence Hall/Dormitory are
# deliberately NOT folded into warehouse/distribution or multifamily -- each
# is its own real-estate asset class with a different mechanical profile, and
# the user's own list named specific categories, not "anything similar."
#
# Multifamily Housing removed 2026-09-07 -- not a default type, still
# selectable explicitly (`property_type=Multifamily Housing` bypasses this
# allowlist the same way any other explicit type does).
DMG_RELEVANT_PROPERTY_TYPES = {
    "Office",                                  # office
    "Medical Office",                          # medical office
    "Hospital (General Medical & Surgical)",   # hospital
    "Other - Specialty Hospital",              # hospital
    "K-12 School",                             # K-12 school
    "College/University",                      # college/university
    "Laboratory",                              # laboratory
    "Data Center",                             # data center
    "Distribution Center",                     # warehouse/distribution
    "Non-Refrigerated Warehouse",              # warehouse/distribution
    "Refrigerated Warehouse",                  # warehouse/distribution
    "Manufacturing/Industrial Plant",          # manufacturing/industrial
    "Hotel",                                   # hotel
    "Retail Store",                            # retail
    "Strip Mall",                              # retail
    "Enclosed Mall",                           # retail
    "Supermarket/Grocery Store",               # retail
    "Wholesale Club/Supercenter",              # retail
    "Lifestyle Center",                        # retail
}

# Word-boundary matched, case-insensitive, plural forms included ("Kaiser
# Foundation Hospitals", "Los Angeles Unified School District") -- "City of"
# is the one multi-word phrase, the rest are single tokens. Deliberately
# \b-bounded: a raw substring check for "inc" matches "distinct"; this doesn't.
_ORG_NAME_HINT_PATTERNS = [
    re.compile(r"\bllc\b", re.I), re.compile(r"\binc\b", re.I),
    re.compile(r"\bhospitals?\b", re.I), re.compile(r"\buniversit(?:y|ies)\b", re.I),
    re.compile(r"\bdistricts?\b", re.I), re.compile(r"\bcount(?:y|ies)\b", re.I),
    re.compile(r"\bcity of\b", re.I), re.compile(r"\bchurch(?:es)?\b", re.I),
    re.compile(r"\bschools?\b", re.I),
]


def looks_like_organization_name(name: str | None) -> bool:
    """True if `name` reads like an organization, not a street address or a
    generic building nickname -- see rank_in_territory's `name_hint`. A hint
    that a rep can call this name, never a claim that it IS the owner: see
    Ab802Building's own docstring on why no owner field exists at all."""
    if not name:
        return False
    return any(p.search(name) for p in _ORG_NAME_HINT_PATTERNS)


def owner_hint_for_name(cfg: Config, name: str | None) -> tuple[str | None, str | None]:
    """(owner, matched_token) for the FIRST entry in config.yaml's
    ab802_owner_hints.tokens that matches Property Name, or (None, None) --
    see that config comment for the editable-without-a-deploy list and
    app/assumptions.py's "AB 802 owner hints" entry for the one regex
    entry's verification (a Prologis internal facility-code prefix,
    checked against Prologis's own public property search before this
    shipped).

    A HINT only, same discipline as looks_like_organization_name: matching
    a token is not a claim that company owns the building, just a name
    worth checking before calling -- see rank_in_territory's own docstring
    and Ab802Building's on why no owner field exists at all.

    Each entry is either `contains` (plain substring, case-insensitive --
    same matching app.call_target._matches_standards_owner already uses
    for its own config-driven owner list) or `regex` (used as written,
    case-insensitive); first match in config list order wins. `matched_token`
    is the entry's own `contains`/`regex` value, for reporting how many
    rows matched EACH configured rule, not just the resulting owner name
    (two different rules can map to the same owner, e.g. the Prologis
    code pattern and the literal "Prologis" substring)."""
    if not name:
        return None, None
    for entry in cfg.get("ab802_owner_hints.tokens") or []:
        owner = entry.get("owner")
        if not owner:
            continue
        pattern = entry.get("regex")
        if pattern:
            if re.match(pattern, name.strip(), re.I):
                return owner, pattern
            continue
        token = entry.get("contains")
        if token and token.lower() in name.lower():
            return owner, token
    return None, None


# AB 802 itself only covers buildings >=50,000 sqft (Public Resources Code
# section 25402.10(a)(1)(A)) -- a row below GFA_FLOOR_SQFT is presumed a data
# error in the source file (the un-floored top 20, measured 2026-09-07,
# included a "Manufacturing/Industrial Plant" row with GFA=100 and a Site EUI
# of 47,628 -- a mis-recorded denominator, not a real building), not a real
# building this program covers. 20,000 rather than the statute's own 50,000
# is deliberate slack: this app has no independent way to confirm GFA against
# an assessor record for most of the state, so the floor only needs to catch
# implausible values, not enforce the statute's own boundary at the margin.
GFA_FLOOR_SQFT = 20_000

# Above this, a row's true eui_ratio is presumed more likely a data error
# than a genuine 5x-median energy hog -- pulled into `anomalies` instead of
# `ranked`, never silently scored. See rank_in_territory's own docstring.
EUI_RATIO_ANOMALY_ABOVE = 5.0


def rank_in_territory(session: Session, cfg: Config, *, county: str | None = None,
                      property_type: str | None = None, year_built_before: int | None = None,
                      eui_above_median: bool = False, has_assessor_match: bool = False,
                      restrict_to_relevant_types: bool = False) -> dict:
    """Returns {"ranked": [...], "anomalies": [...]} -- two terms, summed,
    both shown as plain dict fields rather than folded into one opaque
    score, for every row in `ranked`:

      - age_credit(row.year_built): see that function's own docstring for
        the exact trapezoid (full credit 1985-2012, tapering to 0 by 1975
        and by 2018, 0 before 1975).
      - eui_ratio: row.weather_normalized_site_eui / the MEDIAN weather-
        normalized site EUI for its own primary_property_type, among
        in-territory buildings at or above GFA_FLOOR_SQFT -- relative
        magnitude, not a percentile rank, so a building at 3x its type's
        median EUI clearly outranks one at 1.2x even if both are "top of
        their type." A 1980 office and a 1980 distribution warehouse have
        different normal EUI ranges, so this is always computed within
        type, never board-wide.

    rank_key = eui_ratio + age_credit (0.0 substituted only for summing,
    never displayed as if it were a real 0 -- the caller shows the actual
    per-row field, which stays None when its input was missing). Ties on
    rank_key break on property_gfa_sqft, larger first ("a bigger building
    is a bigger order").

    GFA_FLOOR_SQFT is applied FIRST, to the whole in-territory population,
    before medians are computed -- a single implausible GFA must not be
    allowed to distort the median every other row in its type is ranked
    against. A row below the floor (or with no GFA at all -- unconfirmed,
    never assumed to pass) is dropped entirely, not shown as an anomaly:
    "data error, doesn't belong here" is a different claim than "genuine
    but extreme, verify before calling."

    A row whose eui_ratio EXCEEDS EUI_RATIO_ANOMALY_ABOVE is the second,
    distinct case -- plausibly a real building, but a >5x-median site EUI
    is more often a meter/unit-mismatch or similar data issue than a
    genuine finding. That row is moved to `anomalies` (rank_key is None,
    never sorted into `ranked`) instead of being allowed to dominate the
    top of the list on a number that likely isn't real; the caller labels
    it "data anomaly, verify before calling."

    Median and the EUI-above-median filter are computed over the type's
    FULL floored, in-territory population, before any filter below
    narrows which rows are actually shown -- a filter must never change
    what "median for its type" means for the rows that remain.

    restrict_to_relevant_types (the board's own default -- see
    app/web/main.py:replacement_leads_view): limits results to
    DMG_RELEVANT_PROPERTY_TYPES, ignored when `property_type` names one
    specific type explicitly (a rep who explicitly asks for "Casino" or
    "Multifamily Housing" gets it, default or not).

    Owner-facing fields, most specific first -- the caller shows exactly
    ONE of these per row, never more than one, and shows "no owner data
    available" only when all three are absent:
      1. benchmarking_filer (unchanged, an actual EBEWE filing).
      2. owner_hint/owner_hint_token: see owner_hint_for_name -- a
         config.yaml-driven token/pattern match against Property Name,
         computed only when benchmarking_filer is blank. Labeled
         "owner hint: {owner_hint}" by the caller, never "owner".
      3. name_hint: when BOTH benchmarking_filer and owner_hint are
         blank AND property_name itself reads like an organization
         (looks_like_organization_name), the name is surfaced as a lead
         -- labeled "name on filing" by the caller, never "owner"."""
    rows = [r for r in latest_in_territory_rows(session)
           if r.property_gfa_sqft is not None and r.property_gfa_sqft >= GFA_FLOOR_SQFT]

    by_type: dict[str, list[Ab802Building]] = {}
    for r in rows:
        by_type.setdefault(r.primary_property_type or "Unknown", []).append(r)

    eui_median_by_type: dict[str, float] = {}
    for ptype, group in by_type.items():
        eui_values = [r.weather_normalized_site_eui for r in group
                     if r.weather_normalized_site_eui is not None]
        if eui_values:
            eui_median_by_type[ptype] = statistics.median(eui_values)

    ranked, anomalies = [], []
    for r in rows:
        if county and r.county_from_geocoding != county:
            continue
        if property_type:
            if (r.primary_property_type or "Unknown") != property_type:
                continue
        elif restrict_to_relevant_types and r.primary_property_type not in DMG_RELEVANT_PROPERTY_TYPES:
            continue
        if year_built_before is not None and (r.year_built is None or r.year_built >= year_built_before):
            continue
        median = eui_median_by_type.get(r.primary_property_type or "Unknown")
        if eui_above_median:
            if median is None or r.weather_normalized_site_eui is None or r.weather_normalized_site_eui <= median:
                continue
        if has_assessor_match and not r.assessor_match_method:
            continue

        age = age_credit(r.year_built)
        eui_ratio = (r.weather_normalized_site_eui / median
                    if median and r.weather_normalized_site_eui is not None else None)
        owner_hint = owner_hint_token = None
        if not r.benchmarking_filer:
            owner_hint, owner_hint_token = owner_hint_for_name(cfg, r.property_name)
        name_hint = (r.property_name if not r.benchmarking_filer and not owner_hint
                    and looks_like_organization_name(r.property_name) else None)
        entry = {
            "row": r, "age_credit": age, "eui_ratio": eui_ratio,
            "type_median_eui": median, "name_hint": name_hint,
            "owner_hint": owner_hint, "owner_hint_token": owner_hint_token,
        }
        if eui_ratio is not None and eui_ratio > EUI_RATIO_ANOMALY_ABOVE:
            entry["rank_key"] = None
            anomalies.append(entry)
        else:
            entry["rank_key"] = (eui_ratio or 0.0) + (age or 0.0)
            ranked.append(entry)

    ranked.sort(key=lambda d: (d["rank_key"], d["row"].property_gfa_sqft or 0.0), reverse=True)
    anomalies.sort(key=lambda d: d["eui_ratio"], reverse=True)
    return {"ranked": ranked, "anomalies": anomalies}
