"""CEC IEPR large-load forecast import — a county-level FORWARD MW LAYER, not
a project detector.

Nevada has PUCN dockets (app/sources/pucn.py); California utilities instead
disclose their interconnection request backlog as part of the CEC's biennial
Integrated Energy Policy Report proceeding. SCE's most recent filing
(docket 25-IEPR-03, TN 266008) is a spreadsheet of individual interconnection
requests — but redacted to status / city / region / requested MW / requested
energization year. No developer name, no address, nothing a lead could be
built from. That is deliberate on the utility's part, and it means this data
can never enter the project pipeline the way a CEQA filing or a PUCN docket
does — there is no project to extract, resolve, or attribute a contact to.

What it IS good for: knowing how much large-load MW is queued up, county by
county, ahead of anything showing up in a land-use filing. That is exactly
what Phase 4's spillover score needs (recent/queued data-center MW near a
project raises that project's score) — so this module writes straight into
IeprForwardLoad, a separate table Phase 4 reads directly, never into Signal
or Project.

Manual, twice a year, on purpose. The filing is a snapshot users download
from the CEC's e-filing system by hand (a TN number is not a stable feed
endpoint, and the docket comparison used to derive the churn discount below
was itself a manual read of two snapshots) — building a scraper against it
would imply a cadence and reliability neither the filing nor the docket
process actually has. `scout import-iepr <path.xlsx>` is the entire feed.

Churn discount: comparing SCE's 2024 and 2025 IEPR filings, ~44% of tracked
interconnection requests did not survive unchanged year over year (canceled,
or otherwise dropped from the active list) — an interconnection request is
not a project, and this table's MW must never be read at face value.
`iepr.churn_discount` in config.yaml carries that as an explicit multiplier
applied wherever this layer feeds a score, not baked into the import.
"""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
from sqlmodel import delete, select

from app.models import IeprForwardLoad, utcnow

log = logging.getLogger(__name__)

# City -> county, scoped deliberately to the confirmed cities seen in real SCE
# filings rather than a general CA gazetteer. Two aliases below (Redondo ->
# Redondo Beach, Monterrey Park -> Monterey Park) are included ONLY because
# both the short/misspelled form and the correct form appear as separate rows
# in the same filing, which makes them a data-entry inconsistency rather than
# a guess. Anything else unrecognized ("Open" — a literal placeholder in
# SCE's own sheet, "Near Downtown LA", or a city not yet seen) resolves to
# None rather than being guessed at: null over inference applies to this
# lookup exactly as it does to state and county elsewhere in this codebase.
_CITY_ALIASES = {
    "redondo": "redondo beach",
    "monterrey park": "monterey park",
}
_CITY_TO_COUNTY = {
    "apple valley": "San Bernardino",
    "bloomington": "San Bernardino",
    "carson": "Los Angeles",
    "city of industry": "Los Angeles",
    "culver city": "Los Angeles",
    "delano": "Kern",
    "el segundo": "Los Angeles",
    "goleta": "Santa Barbara",
    "hawthorne": "Los Angeles",
    "huntington park": "Los Angeles",
    "irvine": "Orange",
    "lake elsinore": "Riverside",
    "lake forest": "Orange",
    "lancaster": "Los Angeles",
    "long beach": "Los Angeles",
    "los angeles": "Los Angeles",
    "mojave": "Kern",
    "monterey park": "Los Angeles",
    "moorpark": "Ventura",
    "palm springs": "Riverside",
    "paramount": "Los Angeles",
    "perris": "Riverside",
    "redlands": "San Bernardino",
    "redondo beach": "Los Angeles",
    "san bernardino": "San Bernardino",
    "santa fe springs": "Los Angeles",
    "thousand oaks": "Ventura",
    "trona": "San Bernardino",
    "tulare": "Tulare",
    "tustin": "Orange",
    "vernon": "Los Angeles",
    "victorville": "San Bernardino",
    "west covina": "Los Angeles",
    "whittier": "Los Angeles",
}

# Header text -> field name. Matched case/whitespace-insensitively so a minor
# column-title change ("Requested Peak MW " with its trailing space, seen in
# the real filing) doesn't silently misalign every column after it.
_COLUMNS = {
    "status": "status",
    "cec grouping": "cec_grouping",
    "region": "region",
    "voltage": "voltage",
    "city": "city",
    "requested energization year": "requested_energization_year",
    "requested peak mw": "requested_peak_mw",
}


def _city_to_county(city: str | None) -> str | None:
    if not city:
        return None
    key = city.strip().lower()
    key = _CITY_ALIASES.get(key, key)
    return _CITY_TO_COUNTY.get(key)


def import_iepr_workbook(session, path: str | Path, *, utility: str, docket_tn: str,
                         source_url: str) -> dict:
    """Import one utility's IEPR forecast snapshot. Replaces every prior row
    for this (utility, docket_tn) — see module docstring: the filing is a
    point-in-time snapshot, and re-running the same docket_tn means "I
    re-downloaded this filing," not "here are more rows on top of the old
    ones." Returns import stats, including unmapped cities, so a bad column
    read or an unrecognized new city shows up immediately rather than as
    silently-null counties downstream in Phase 4.
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.worksheets[0]

    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_index: dict[str, int] = {}
    for i, raw in enumerate(header_row):
        if not raw:
            continue
        key = " ".join(str(raw).split()).lower()
        field = _COLUMNS.get(key)
        if field:
            col_index[field] = i

    missing = [f for f in _COLUMNS.values() if f not in col_index]
    if missing:
        raise ValueError(
            f"IEPR workbook is missing expected column(s): {missing} — "
            f"header row was {header_row!r}; this importer assumes the same "
            f"layout as TN 266008 and needs a look if the utility changed it.")

    # Delete the prior snapshot for this (utility, docket_tn) BEFORE adding the
    # new rows below — both inside the caller's transaction, so a failed
    # import (a raised exception mid-parse) rolls back to the prior snapshot
    # rather than leaving the table empty.
    session.exec(delete(IeprForwardLoad).where(
        IeprForwardLoad.utility == utility, IeprForwardLoad.docket_tn == docket_tn))

    rows_in = 0
    rows_stored = 0
    unmapped_cities: dict[str, int] = {}
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in raw_row):
            continue
        rows_in += 1

        def cell(field: str):
            return raw_row[col_index[field]]

        city = cell("city")
        county = _city_to_county(city)
        if city and county is None:
            unmapped_cities[city] = unmapped_cities.get(city, 0) + 1

        year_raw = cell("requested_energization_year")
        mw_raw = cell("requested_peak_mw")
        voltage_raw = cell("voltage")
        grouping = cell("cec_grouping")

        session.add(IeprForwardLoad(
            utility=utility,
            docket_tn=docket_tn,
            source_url=source_url,
            status=str(cell("status")).strip() if cell("status") else None,
            cec_grouping=str(grouping).strip() if grouping is not None else None,
            region=str(cell("region")).strip() if cell("region") else None,
            city=str(city).strip() if city else None,
            county=county,
            voltage=float(voltage_raw) if isinstance(voltage_raw, (int, float)) else None,
            requested_energization_year=int(year_raw) if isinstance(year_raw, (int, float)) else None,
            requested_peak_mw=float(mw_raw) if isinstance(mw_raw, (int, float)) else None,
        ))
        rows_stored += 1

    session.flush()

    return {
        "utility": utility,
        "docket_tn": docket_tn,
        "rows_in_sheet": rows_in,
        "rows_stored": rows_stored,
        "unmapped_cities": unmapped_cities,
        "imported_at": utcnow().isoformat(),
    }


def county_forward_mw(session, *, utility: str | None = None,
                      exclude_statuses: tuple[str, ...] = ("Canceled", "Distribution-Canceled")) -> dict[str, float]:
    """County -> total requested MW, summed across every stored IEPR row not
    in an excluded (dead) status. Rows with no resolvable county are excluded
    from the total rather than silently attributed anywhere — see
    _city_to_county. Callers apply config's iepr.churn_discount themselves;
    this function returns the raw filed figure so the discount stays visible
    and tunable at the call site (Phase 4 scoring) instead of buried here.
    """
    stmt = select(IeprForwardLoad).where(IeprForwardLoad.county.is_not(None))
    if utility:
        stmt = stmt.where(IeprForwardLoad.utility == utility)
    rows = session.exec(stmt).all()
    totals: dict[str, float] = {}
    for row in rows:
        if row.status in exclude_statuses:
            continue
        if row.requested_peak_mw is None:
            continue
        totals[row.county] = totals.get(row.county, 0.0) + row.requested_peak_mw
    return totals
